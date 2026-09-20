import re
import json
import os
import sys
import math
import numpy as np
from collections import defaultdict, deque
from datetime import datetime

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.backends.backend_svg import FigureCanvasSVG

import ezdxf
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

STEP_FILEPATH = r'c:\Users\yh\.trae-cn\attachments\6a9698e44f0e5b6cccc8a92f\e310c563-5814-4a8b-8303-751e1e4612dc_f9f74891-2cd8-41c9-9c67-11a4c4c66e78_舱体+基座模型.stp'
PROJECT_DIR = r'C:\Users\yh\AppData\Roaming\TRAE SOLO CN\ModularData\ai-agent\work-mode-projects\6a9698e44f0e5b6cccc8a92c'
OUTPUT_DIR = os.path.join(PROJECT_DIR, 'production_files')
SVG_DIR = os.path.join(OUTPUT_DIR, 'SVG')
DXF_DIR = os.path.join(OUTPUT_DIR, 'DXF')
STEP_DIR = os.path.join(OUTPUT_DIR, 'STEP')
BOM_DIR = os.path.join(OUTPUT_DIR, 'BOM')

for d in [OUTPUT_DIR, SVG_DIR, DXF_DIR, STEP_DIR, BOM_DIR]:
    os.makedirs(d, exist_ok=True)

# ============================================================
# STEP File Parser (fixed regex)
# ============================================================
def parse_step_file(filepath):
    entities = {}
    entity_order = []
    with open(filepath, 'r', encoding='utf-8', errors='replace') as f:
        in_data = False
        for line in f:
            line = line.strip()
            if not in_data:
                if line == 'DATA;':
                    in_data = True
                continue
            if line == 'ENDSEC;':
                break
            if not line or line.startswith('//') or line.startswith('/*'):
                continue
            match = re.match(r'#(\d+)=(.*)', line)
            if match:
                eid = int(match.group(1))
                rest = match.group(2)
                while not rest.rstrip().endswith(';'):
                    next_line = f.readline()
                    if not next_line:
                        break
                    rest += ' ' + next_line.strip()
                rest = rest.rstrip()
                if rest.endswith(';'):
                    rest = rest[:-1]
                if rest.startswith('('):
                    inner = rest[1:]
                    type_match = re.match(r'([A-Z][A-Z0-9_]*)\s*\(', inner)
                    if type_match:
                        etype = type_match.group(1)
                        if 'REPRESENTATION_RELATIONSHIP_WITH_TRANSFORMATION' in rest:
                            etype = 'REPRESENTATION_RELATIONSHIP_WITH_TRANSFORMATION'
                        elif 'SHAPE_REPRESENTATION_RELATIONSHIP' in rest:
                            etype = 'SHAPE_REPRESENTATION_RELATIONSHIP'
                        entities[eid] = (etype, rest)
                        entity_order.append(eid)
                    else:
                        entities[eid] = ('COMPOUND', rest)
                        entity_order.append(eid)
                else:
                    type_match = re.match(r'([A-Z][A-Z0-9_]*)\s*\(', rest)
                    if type_match:
                        etype = type_match.group(1)
                        entities[eid] = (etype, rest)
                        entity_order.append(eid)
    return entities, entity_order

def extract_refs(value_str):
    return [int(r) for r in re.findall(r'#(\d+)', value_str)]

def get_param_list(entity_str):
    paren_start = entity_str.find('(')
    if paren_start == -1:
        return []
    depth = 0
    last_paren = -1
    for i in range(paren_start, len(entity_str)):
        c = entity_str[i]
        if c == '(':
            depth += 1
        elif c == ')':
            depth -= 1
            if depth == 0:
                last_paren = i
                break
    if last_paren == -1:
        return []
    inner = entity_str[paren_start+1:last_paren]
    return split_params(inner)

def split_params(s):
    params = []
    depth = 0
    current = ''
    in_string = False
    for c in s:
        if c == "'" and not in_string:
            in_string = True
            current += c
        elif c == "'" and in_string:
            in_string = False
            current += c
        elif c == '(' and not in_string:
            depth += 1
            current += c
        elif c == ')' and not in_string:
            depth -= 1
            current += c
        elif c == ',' and depth == 0 and not in_string:
            params.append(current.strip())
            current = ''
        else:
            current += c
    if current.strip():
        params.append(current.strip())
    return params

def eval_exponent(s):
    s = s.strip()
    parts = s.split('E')
    if len(parts) == 2:
        base = float(parts[0])
        exp = int(parts[1])
        return base * (10 ** exp)
    return float(s)

def parse_cartesian_point(entity_str):
    params = get_param_list(entity_str)
    for p in reversed(params):
        p = p.strip()
        if p.startswith('(') and p.endswith(')'):
            inner = p[1:-1]
            if ',' in inner:
                values = inner.split(',')
                coords = []
                for v in values:
                    try:
                        coords.append(eval_exponent(v.strip()))
                    except:
                        pass
                if len(coords) >= 3:
                    return coords[:3]
    return None

def parse_direction(entity_str):
    params = get_param_list(entity_str)
    for p in reversed(params):
        p = p.strip()
        if p.startswith('(') and p.endswith(')'):
            inner = p[1:-1]
            if ',' in inner:
                values = inner.split(',')
                coords = []
                for v in values:
                    try:
                        coords.append(eval_exponent(v.strip()))
                    except:
                        pass
                if len(coords) >= 3:
                    return coords[:3]
    return [0, 0, 1]

# ============================================================
# Geometry Extraction
# ============================================================
def extract_part_geometry(entities, brep_id):
    edges = []
    visited = set()

    def get_point(eid):
        entity = entities.get(eid)
        if not entity or entity[0] != 'CARTESIAN_POINT':
            return None
        return parse_cartesian_point(entity[1])

    def process_edge_curve(eid):
        if eid in visited:
            return
        visited.add(eid)
        entity = entities.get(eid)
        if not entity:
            return
        refs = extract_refs(entity[1])
        v1, v2 = None, None
        for r in refs:
            rtype = entities.get(r, (None,))[0]
            if rtype == 'VERTEX_POINT':
                vrefs = extract_refs(entities[r][1])
                for vr in vrefs:
                    if entities.get(vr, (None,))[0] == 'CARTESIAN_POINT':
                        pt = get_point(vr)
                        if pt:
                            if v1 is None:
                                v1 = pt
                            else:
                                v2 = pt
            elif rtype == 'LINE':
                line_refs = extract_refs(entities[r][1])
                for lr in line_refs:
                    if entities.get(lr, (None,))[0] == 'CARTESIAN_POINT':
                        pt = get_point(lr)
                        if pt:
                            if v1 is None:
                                v1 = pt
                            elif v2 is None:
                                v2 = pt
                    elif entities.get(lr, (None,))[0] == 'VECTOR':
                        vec_refs = extract_refs(entities[lr][1])
                        for vr in vec_refs:
                            if entities.get(vr, (None,))[0] == 'DIRECTION':
                                d = parse_direction(entities[vr][1])
                                vec_entity = entities.get(lr)
                                if vec_entity:
                                    vparams = get_param_list(vec_entity[1])
                                    mag = 0.0
                                    if len(vparams) > 1:
                                        try:
                                            mag = float(vparams[1].strip().strip("'"))
                                        except:
                                            pass
                                    if v1 and d:
                                        v2 = [v1[i] + d[i] * mag for i in range(3)]
            elif rtype == 'CIRCLE':
                circle_refs = extract_refs(entities[r][1])
                center = None
                radius = 0
                axis_dir = [0, 0, 1]
                for cr in circle_refs:
                    crtype = entities.get(cr, (None,))[0]
                    if crtype == 'CARTESIAN_POINT':
                        center = get_point(cr)
                    elif crtype == 'DIRECTION':
                        axis_dir = parse_direction(entities[cr][1])
                circle_params = get_param_list(entities[r][1])
                if len(circle_params) > 2:
                    try:
                        radius = float(circle_params[2].strip().strip("'"))
                    except:
                        pass
                elif len(circle_params) > 1:
                    try:
                        radius = float(circle_params[1].strip().strip("'"))
                    except:
                        pass
                if center and radius > 0:
                    theta = np.linspace(0, 2*np.pi, 30)
                    if abs(axis_dir[2]) > 0.9:
                        xs = center[0] + radius * np.cos(theta)
                        ys = center[1] + radius * np.sin(theta)
                        zs = np.full_like(theta, center[2])
                    elif abs(axis_dir[1]) > 0.9:
                        xs = center[0] + radius * np.cos(theta)
                        ys = np.full_like(theta, center[1])
                        zs = center[2] + radius * np.sin(theta)
                    else:
                        xs = np.full_like(theta, center[0])
                        ys = center[1] + radius * np.cos(theta)
                        zs = center[2] + radius * np.sin(theta)
                    for i in range(len(theta)-1):
                        edges.append(([xs[i], ys[i], zs[i]], [xs[i+1], ys[i+1], zs[i+1]]))
        if v1 and v2:
            edges.append((v1, v2))

    def process_oriented_edge(eid):
        if eid in visited:
            return
        visited.add(eid)
        entity = entities.get(eid)
        if not entity:
            return
        refs = extract_refs(entity[1])
        for r in refs:
            if entities.get(r, (None,))[0] == 'EDGE_CURVE':
                process_edge_curve(r)

    def process_edge_loop(eid):
        if eid in visited:
            return
        visited.add(eid)
        entity = entities.get(eid)
        if not entity:
            return
        refs = extract_refs(entity[1])
        for r in refs:
            rtype = entities.get(r, (None,))[0]
            if rtype == 'ORIENTED_EDGE':
                process_oriented_edge(r)
            elif rtype == 'EDGE_LOOP':
                process_edge_loop(r)

    def process_face_bound(eid):
        if eid in visited:
            return
        visited.add(eid)
        entity = entities.get(eid)
        if not entity:
            return
        refs = extract_refs(entity[1])
        for r in refs:
            rtype = entities.get(r, (None,))[0]
            if rtype in ('EDGE_LOOP', 'VERTEX_LOOP'):
                process_edge_loop(r)

    def process_advanced_face(eid):
        if eid in visited:
            return
        visited.add(eid)
        entity = entities.get(eid)
        if not entity:
            return
        refs = extract_refs(entity[1])
        for r in refs:
            rtype = entities.get(r, (None,))[0]
            if rtype in ('FACE_BOUND', 'FACE_OUTER_BOUND'):
                process_face_bound(r)

    def process_shell(eid):
        if eid in visited:
            return
        visited.add(eid)
        entity = entities.get(eid)
        if not entity:
            return
        refs = extract_refs(entity[1])
        for r in refs:
            if entities.get(r, (None,))[0] == 'ADVANCED_FACE':
                process_advanced_face(r)

    def process_brep(eid):
        if eid in visited:
            return
        visited.add(eid)
        entity = entities.get(eid)
        if not entity:
            return
        refs = extract_refs(entity[1])
        for r in refs:
            rtype = entities.get(r, (None,))[0]
            if rtype in ('CLOSED_SHELL', 'OPEN_SHELL', 'ORIENTED_CLOSED_SHELL'):
                process_shell(r)

    brep_entity = entities.get(brep_id)
    if brep_entity:
        refs = extract_refs(brep_entity[1])
        for r in refs:
            rtype = entities.get(r, (None,))[0]
            if rtype == 'MANIFOLD_SOLID_BREP':
                process_brep(r)

    return edges

# ============================================================
# STEP Single Part Export
# ============================================================
def collect_entity_ids(entities, start_id):
    collected = set()
    queue = deque([start_id])
    while queue:
        eid = queue.popleft()
        if eid in collected:
            continue
        collected.add(eid)
        entity = entities.get(eid)
        if entity:
            refs = extract_refs(entity[1])
            for r in refs:
                if r not in collected and r in entities:
                    queue.append(r)
    return collected

def write_step_file(entities, entity_ids, output_path, part_name):
    sorted_ids = sorted(entity_ids)
    lines = []
    lines.append('ISO-10303-21;')
    lines.append('HEADER;')
    lines.append(f"FILE_DESCRIPTION(('Part: {part_name}'), '2;1');")
    lines.append(f"FILE_NAME('{part_name}.stp','{datetime.now().strftime('%Y-%m-%dT%H:%M:%S')}',(''),(''),'TRADE_AGENT','','');")
    lines.append("FILE_SCHEMA(('AP214_ARM'));")
    lines.append('ENDSEC;')
    lines.append('DATA;')
    id_map = {}
    new_id = 1
    for old_id in sorted_ids:
        id_map[old_id] = new_id
        new_id += 1
    for old_id in sorted_ids:
        entity = entities.get(old_id)
        if entity:
            estr = entity[1]
            def replace_ref(m):
                ref_id = int(m.group(1))
                if ref_id in id_map:
                    return f'#{id_map[ref_id]}'
                return m.group(0)
            estr = re.sub(r'#(\d+)', replace_ref, estr)
            lines.append(f'#{id_map[old_id]}={estr};')
    lines.append('ENDSEC;')
    lines.append('END-ISO-10303-21;')
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write('\n'.join(lines))

# ============================================================
# SVG 2D Drawing
# ============================================================
def generate_svg_drawing(edges, bbox, part_name, output_path):
    if not edges:
        return False
    all_pts = np.array([p for e in edges for p in e])
    mn = all_pts.min(axis=0)
    mx = all_pts.max(axis=0)
    dx, dy, dz = mx - mn
    size = max(dx, dy, dz)
    if size < 0.001:
        size = 1

    fig, axes = plt.subplots(2, 2, figsize=(14, 12), facecolor='white')
    fig.suptitle(f'{part_name} - 2D Engineering Drawing', fontsize=14, fontweight='bold', y=0.98)

    views = [
        ('Top View (XY)', 0, 1, axes[0, 0]),
        ('Front View (XZ)', 0, 2, axes[0, 1]),
        ('Side View (YZ)', 1, 2, axes[1, 0]),
    ]

    for title, ai, bi, ax in views:
        ax.set_title(title, fontsize=11, fontweight='bold')
        ax.set_aspect('equal')
        ax.set_facecolor('#fafafa')

        for e in edges:
            x1, y1 = e[0][ai], e[0][bi]
            x2, y2 = e[1][ai], e[1][bi]
            ax.plot([x1, x2], [y1, y2], 'b-', linewidth=0.5, alpha=0.7)

        v_mn = [mn[ai], mn[bi]]
        v_mx = [mx[ai], mx[bi]]
        v_dx = v_mx[0] - v_mn[0]
        v_dy = v_mx[1] - v_mn[1]

        margin = max(v_dx, v_dy) * 0.1
        ax.set_xlim(v_mn[0] - margin, v_mx[0] + margin)
        ax.set_ylim(v_mn[1] - margin, v_mx[1] + margin)

        ax.plot([v_mn[0], v_mx[0], v_mx[0], v_mn[0], v_mn[0]],
                [v_mn[1], v_mn[1], v_mx[1], v_mx[1], v_mn[1]],
                'r--', linewidth=1.5, alpha=0.5)

        ax.annotate(f'{v_dx:.1f}', xy=((v_mn[0]+v_mx[0])/2, v_mn[1]-margin*0.5),
                    ha='center', va='top', fontsize=9, color='red', fontweight='bold')
        ax.annotate(f'{v_dy:.1f}', xy=(v_mx[0]+margin*0.3, (v_mn[1]+v_mx[1])/2),
                    ha='left', va='center', fontsize=9, color='green', fontweight='bold')

        ax.set_xlabel(f'{["X","Y","Z"][ai]} (mm)', fontsize=9)
        ax.set_ylabel(f'{["X","Y","Z"][bi]} (mm)', fontsize=9)
        ax.grid(True, alpha=0.3)
        ax.tick_params(labelsize=7)

    info_ax = axes[1, 1]
    info_ax.axis('off')
    info = (
        f'Part Name: {part_name}\n\n'
        f'Dimensions:\n'
        f'  X = {dx:.1f} mm\n'
        f'  Y = {dy:.1f} mm\n'
        f'  Z = {dz:.1f} mm\n\n'
        f'Bounding Box:\n'
        f'  Min: ({mn[0]:.1f}, {mn[1]:.1f}, {mn[2]:.1f})\n'
        f'  Max: ({mx[0]:.1f}, {mx[1]:.1f}, {mx[2]:.1f})\n\n'
        f'Edges: {len(edges)}\n'
        f'Points: {len(all_pts)}'
    )
    info_ax.text(0.1, 0.9, info, transform=info_ax.transAxes,
                 fontsize=10, verticalalignment='top', fontfamily='monospace',
                 bbox=dict(boxstyle='round', facecolor='wheat', alpha=0.5))

    plt.tight_layout(rect=[0, 0, 1, 0.96])
    fig.savefig(output_path, format='svg', bbox_inches='tight', facecolor='white')
    plt.close(fig)
    return True

# ============================================================
# DXF 2D Drawing
# ============================================================
def generate_dxf_drawing(edges, bbox, part_name, output_path):
    if not edges:
        return False
    all_pts = np.array([p for e in edges for p in e])
    mn = all_pts.min(axis=0)
    mx = all_pts.max(axis=0)
    dx, dy, dz = mx - mn

    doc = ezdxf.new('R2010', setup=True)
    doc.layers.add('OUTLINE', color=5)
    doc.layers.add('DIMENSIONS', color=1)
    doc.layers.add('BBOX', color=2)
    doc.layers.add('TEXT', color=7)

    msp = doc.modelspace()

    scale = 0.001

    views = [
        ('TOP', 0, 1, 0),
        ('FRONT', 0, 2, 1),
        ('SIDE', 1, 2, 2),
    ]

    view_offset_x = 0
    for view_name, ai, bi, vidx in views:
        offset_x = view_offset_x
        offset_y = vidx * (max(mx - mn) * scale + 50)

        for e in edges:
            x1 = e[0][ai] * scale + offset_x
            y1 = e[0][bi] * scale + offset_y
            x2 = e[1][ai] * scale + offset_x
            y2 = e[1][bi] * scale + offset_y
            msp.add_line((x1, y1), (x2, y2), dxfattribs={'layer': 'OUTLINE'})

        v_mn = [mn[ai] * scale, mn[bi] * scale]
        v_mx = [mx[ai] * scale, mx[bi] * scale]
        v_dx = v_mx[0] - v_mn[0]
        v_dy = v_mx[1] - v_mn[1]

        corners = [
            (v_mn[0] + offset_x, v_mn[1] + offset_y),
            (v_mx[0] + offset_x, v_mn[1] + offset_y),
            (v_mx[0] + offset_x, v_mx[1] + offset_y),
            (v_mn[0] + offset_x, v_mx[1] + offset_y),
            (v_mn[0] + offset_x, v_mn[1] + offset_y),
        ]
        for i in range(len(corners) - 1):
            msp.add_line(corners[i], corners[i+1], dxfattribs={'layer': 'BBOX'})

        msp.add_text(f'{view_name} VIEW - {part_name}',
                     dxfattribs={'layer': 'TEXT', 'height': 5}).set_placement(
            (offset_x, v_mx[1] + offset_y + 10))

        v_dims = [mx[ai] - mn[ai], mx[bi] - mn[bi]]
        dim_labels = [f'{v_dims[0]:.1f} mm', f'{v_dims[1]:.1f} mm']
        msp.add_text(f'W={dim_labels[0]}  H={dim_labels[1]}',
                     dxfattribs={'layer': 'DIMENSIONS', 'height': 4}).set_placement(
            (offset_x, v_mn[1] + offset_y - 8))

        view_offset_x += v_dx + 20

    doc.saveas(output_path)
    return True

# ============================================================
# BOM Excel
# ============================================================
def generate_bom(parts_info, output_path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'BOM'

    headers = ['No.', 'Part Name', 'X (mm)', 'Y (mm)', 'Z (mm)', 'Volume (mm3)', 'Point Count', 'Note']
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_font = Font(color='FFFFFF', bold=True, size=11)
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = thin_border

    name_counts = defaultdict(int)
    for part in parts_info:
        name_counts[part['name']] += 1

    unique_names = set()
    for idx, part in enumerate(parts_info, 1):
        name = part['name']
        if name in unique_names:
            continue
        unique_names.add(name)
        bbox = part.get('bbox')
        if bbox and bbox.get('size'):
            size = bbox['size']
            volume = size[0] * size[1] * size[2]
            point_count = bbox.get('point_count', 0)
        else:
            size = [0, 0, 0]
            volume = 0
            point_count = 0
        qty = name_counts[name]

        row_data = [idx, name, size[0], size[1], size[2], volume, point_count, '']
        if qty > 1:
            row_data[7] = f'Qty: {qty}'
        for col, val in enumerate(row_data, 1):
            cell = ws.cell(row=len(unique_names) + 1, column=col, value=val)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal='center' if col != 2 else 'left')

    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 35
    ws.column_dimensions['C'].width = 12
    ws.column_dimensions['D'].width = 12
    ws.column_dimensions['E'].width = 12
    ws.column_dimensions['F'].width = 15
    ws.column_dimensions['G'].width = 12
    ws.column_dimensions['H'].width = 12

    ws.auto_filter.ref = f'A1:H{len(unique_names) + 1}'

    ws2 = wb.create_sheet('Summary')
    ws2['A1'] = 'BOM Summary'
    ws2['A1'].font = Font(bold=True, size=14)
    ws2['A3'] = 'Total Parts (instances):'
    ws2['B3'] = len(parts_info)
    ws2['A4'] = 'Unique Parts:'
    ws2['B4'] = len(unique_names)
    ws2['A5'] = 'Parts with Duplicates:'
    ws2['B5'] = sum(1 for n, c in name_counts.items() if c > 1)
    ws2['A6'] = 'Generated Date:'
    ws2['B6'] = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    ws2['A8'] = 'Part Name'
    ws2['B8'] = 'Quantity'
    ws2['A8'].font = Font(bold=True)
    ws2['B8'].font = Font(bold=True)
    for i, (name, count) in enumerate(sorted(name_counts.items(), key=lambda x: -x[1]), 9):
        ws2[f'A{i}'] = name
        ws2[f'B{i}'] = count

    ws2.column_dimensions['A'].width = 35
    ws2.column_dimensions['B'].width = 12

    wb.save(output_path)
    return len(unique_parts_from_info(parts_info))

def unique_parts_from_info(parts_info):
    seen = set()
    result = []
    for p in parts_info:
        if p['name'] not in seen:
            seen.add(p['name'])
            result.append(p)
    return result

# ============================================================
# Main
# ============================================================
def main():
    print("=" * 60)
    print("Production Files Generator")
    print("=" * 60)

    parts_info_path = os.path.join(PROJECT_DIR, 'step_parts_info.json')
    with open(parts_info_path, 'r', encoding='utf-8') as f:
        parts_info = json.load(f)

    unique_parts = unique_parts_from_info(parts_info)
    print(f"Total parts: {len(parts_info)}, Unique: {len(unique_parts)}")

    # 1. Generate BOM
    print("\n[1/4] Generating BOM Excel...")
    bom_path = os.path.join(BOM_DIR, 'BOM.xlsx')
    generate_bom(parts_info, bom_path)
    print(f"  BOM saved: {bom_path}")

    # 2. Parse STEP file for geometry extraction and STEP export
    print("\n[2/4] Parsing STEP file...")
    entities, entity_order = parse_step_file(STEP_FILEPATH)
    print(f"  Parsed {len(entities)} entities")

    # 3. Generate SVG, DXF, and STEP for each unique part
    print(f"\n[3/4] Generating SVG, DXF, and STEP files for {len(unique_parts)} parts...")

    svg_count = 0
    dxf_count = 0
    step_count = 0
    fail_count = 0

    for i, part in enumerate(unique_parts):
        name = part['name']
        brep_id = part['brep_id']
        product_id = part['product_id']
        safe_name = re.sub(r'[^\w\-]', '_', name)
        file_base = f"part_{i+1:04d}_{safe_name}"

        try:
            edges = extract_part_geometry(entities, brep_id)

            if edges:
                svg_path = os.path.join(SVG_DIR, f'{file_base}.svg')
                if generate_svg_drawing(edges, part['bbox'], name, svg_path):
                    svg_count += 1

                dxf_path = os.path.join(DXF_DIR, f'{file_base}.dxf')
                if generate_dxf_drawing(edges, part['bbox'], name, dxf_path):
                    dxf_count += 1
            else:
                print(f"  [{i+1}/{len(unique_parts)}] {name} - no geometry extracted")

            step_ids = collect_entity_ids(entities, brep_id)
            step_path = os.path.join(STEP_DIR, f'{file_base}.stp')
            write_step_file(entities, step_ids, step_path, name)
            step_count += 1

        except Exception as e:
            print(f"  [{i+1}/{len(unique_parts)}] {name} - ERROR: {e}")
            fail_count += 1

        if (i + 1) % 50 == 0:
            print(f"  Progress: {i+1}/{len(unique_parts)} (SVG:{svg_count} DXF:{dxf_count} STEP:{step_count})")

    print(f"\n[4/4] Summary")
    print(f"  SVG files: {svg_count}")
    print(f"  DXF files: {dxf_count}")
    print(f"  STEP files: {step_count}")
    print(f"  Failures: {fail_count}")
    print(f"  BOM: 1 file")
    print(f"\n  Output directory: {OUTPUT_DIR}")
    print(f"  - BOM/BOM.xlsx")
    print(f"  - SVG/ ({svg_count} files)")
    print(f"  - DXF/ ({dxf_count} files)")
    print(f"  - STEP/ ({step_count} files)")
    print("\nDone!")

if __name__ == '__main__':
    main()
